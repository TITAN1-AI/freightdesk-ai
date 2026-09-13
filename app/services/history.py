"""Structured historical staging/query store, separate from current canonical shipments."""
import csv
import hashlib
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from statistics import median
from zipfile import ZipFile

from pydantic import ValidationError

from app.core.runtime import RuntimePaths
from app.models.domain import AuditEvent, Role, utcnow
from app.models.history import (
    ApprovedCustomerRule, DerivedProfile, HistoricalLoad, HistoricalProvenance,
    ProposedCustomerRule, RateHistorySummary,
)
from app.services.mail_sync import digest
from app.services.store import Store

DATES = {"pickup_date", "delivery_date", "created_date"}
MONEY = {"customer_revenue", "carrier_pay", "gross_margin", "weight"}
QUERY_FIELDS = {"customer", "origin", "destination", "equipment", "carrier", "pickup_facility",
                "delivery_facility", "load_number", "carrier_mc", "carrier_dot"}


def history_store():
    return Store(RuntimePaths.from_environment().path("Data", "booking-logistics", "history", "history.sqlite3"))


class AscendHistoricalImporter:
    def __init__(self, store, tenant="booking-logistics"):
        self.store, self.tenant = store, tenant

    @staticmethod
    def read_export(path: Path):
        root = RuntimePaths.from_environment().path("Data", "booking-logistics", "history", "inbox")
        if not path.absolute().is_relative_to(root.absolute()):
            raise PermissionError("Copy the owner-provided export to the approved runtime history inbox")
        relative = path.absolute().relative_to(root.absolute())
        checked = RuntimePaths.from_environment().path(
            "Data", "booking-logistics", "history", "inbox", *relative.parts)
        if checked.stat().st_size > 20_000_000:
            raise ValueError("Export size limit exceeded")
        data = checked.read_bytes()
        source_hash = hashlib.sha256(data).hexdigest()
        if checked.suffix.casefold() == ".csv":
            import io
            reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
            headers = reader.fieldnames or []
            if len(headers) > 100 or len(set(headers)) != len(headers):
                raise ValueError("Export headers invalid or duplicated")
            rows = []
            for row in reader:
                if len(rows) >= 20_000:
                    raise ValueError("Export row limit exceeded")
                rows.append(row)
        elif checked.suffix.casefold() == ".xlsx":
            with ZipFile(checked) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50_000_000:
                    raise ValueError("Expanded workbook exceeds limit")
                if any("vbaProject" in item.filename or "externalLinks/" in item.filename for item in archive.infolist()):
                    raise PermissionError("Workbook macros/external links are not supported")
            from openpyxl import load_workbook
            book = load_workbook(checked, read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) != 1:
                    raise ValueError("Select/export one worksheet explicitly before import")
                sheet = book.worksheets[0]
                if sheet.max_row > 20_001 or sheet.max_column > 100:
                    raise ValueError("Workbook dimensions exceed bound")
                iterator = sheet.iter_rows()
                headers = [str(cell.value or "") for cell in next(iterator)]
                if not all(headers) or len(set(headers)) != len(headers):
                    raise ValueError("Workbook headers invalid")
                rows = []
                for cells in iterator:
                    if any(cell.data_type == "f" for cell in cells):
                        raise PermissionError("Export values only; formulas are not evaluated or trusted")
                    rows.append(dict(zip(headers, [cell.value for cell in cells], strict=True)))
            finally:
                book.close()
        else:
            raise ValueError("Only owner-provided CSV/XLSX value exports are supported")
        return {"headers": headers, "rows": rows, "source_hash": source_hash}

    def stage(self, export, mapping, source_identifier, source_system="ascend_export"):
        if not mapping.reviewed or "load_number" not in mapping.columns:
            raise PermissionError("Inspect actual headers and obtain explicit reviewed schema mapping first")
        if set(mapping.columns) - set(HistoricalLoad.model_fields):
            raise ValueError("Unknown canonical history fields")
        if any(column not in export["headers"] for column in mapping.columns.values()):
            raise ValueError("Mapped export columns are absent")
        batch_id = digest([source_system, export["source_hash"], mapping.model_dump()])
        rows = []
        for index, row in enumerate(export["rows"]):
            values, errors = {}, []
            for field, column in mapping.columns.items():
                raw = row.get(column)
                if raw is None or str(raw).strip() == "":
                    continue
                try:
                    if isinstance(raw, str) and raw.startswith(("=", "+@", "-@", "@")):
                        raise ValueError("formula-like source")
                    if field in DATES:
                        values[field] = (raw.date() if isinstance(raw, datetime) else raw if isinstance(raw, date)
                            else datetime.strptime(str(raw), mapping.date_formats[field]).date()
                            if field in mapping.date_formats else date.fromisoformat(str(raw)))
                    elif field in MONEY:
                        value = Decimal(str(raw).replace(",", "").replace("$", "").strip())
                        if not value.is_finite():
                            raise ValueError("nonfinite value")
                        values[field] = value
                    else:
                        values[field] = str(raw).strip()
                except (ValueError, InvalidOperation, TypeError):
                    errors.append(field)
            if mapping.currency:
                values["currency"] = mapping.currency
            try:
                normalized = HistoricalLoad.model_validate(values)
                if any(getattr(normalized, f) is not None for f in MONEY - {"weight"}) and not normalized.currency:
                    errors.append("currency_required")
                if normalized.customer_revenue is not None and normalized.carrier_pay is not None:
                    computed = normalized.customer_revenue - normalized.carrier_pay
                    if normalized.gross_margin is not None and abs(normalized.gross_margin - computed) > Decimal(".01"):
                        errors.append("gross_margin_conflict")
                payload = normalized.model_dump(mode="json")
            except ValidationError:
                payload = None
                errors.append("normalization_failed")
            rows.append({"row": index + 2, "data": payload, "errors": errors,
                         "status": "INVALID" if errors else "VALID"})
        with self.store.transaction():
            self.store.put(self.tenant, "history_batch", batch_id, {
                "id": batch_id, "source_identifier": source_identifier, "source_system": source_system,
                "source_hash": export["source_hash"], "mapping": mapping.model_dump(), "rows": rows,
                "staged_at": utcnow().isoformat()})
        return {"batch_id": batch_id, "valid": sum(r["status"] == "VALID" for r in rows),
                "invalid": sum(r["status"] == "INVALID" for r in rows)}

    def commit(self, batch_id, actor):
        if actor.tenant_id != self.tenant or actor.role != Role.OWNER:
            raise PermissionError("Owner historical import authorization required")
        inserted, duplicates, conflicts = 0, 0, 0
        with self.store.transaction():
            batch = self.store.get(self.tenant, "history_batch", batch_id)
            existing = self.store.all(self.tenant, "history_record")
            for row in batch["rows"]:
                if row["status"] != "VALID":
                    continue
                data = row["data"]
                logical = [batch["source_system"], data["load_number"]]
                record_id = digest([logical, data])
                if any(record["id"] == record_id for record in existing):
                    duplicates += 1
                    continue
                if any(record["logical_key"] == logical for record in existing):
                    row["status"] = "REVIEW_REQUIRED"
                    row["errors"] = ["existing_historical_version_conflict"]
                    conflicts += 1
                    continue
                provenance = HistoricalProvenance(source_system=batch["source_system"],
                    source_identifier=batch["source_identifier"], source_load_id=data["load_number"],
                    imported_at=utcnow(), source_hash=batch["source_hash"], source_version=record_id,
                    normalization_status="VALID")
                record = {"id": record_id, "logical_key": logical, "data": data,
                          "provenance": provenance.model_dump(mode="json")}
                self.store.put(self.tenant, "history_record", record_id, record)
                existing.append(record)
                inserted += 1
            self.store.put(self.tenant, "history_batch", batch_id, batch)
            self.store.audit(AuditEvent(tenant_id=self.tenant, source="history_import", actor_id=actor.id,
                event="HISTORY_IMPORT_COMMITTED", explanation="Reviewed historical rows; canonical loads unchanged",
                facts={"inserted": inserted, "duplicates": duplicates, "conflicts": conflicts}))
        return {"inserted": inserted, "duplicates": duplicates, "conflicts": conflicts,
                "live_shipments_changed": 0}


class HistoricalIntelligence:
    def __init__(self, store, tenant="booking-logistics"):
        self.store, self.tenant = store, tenant

    def query(self, *, date_from=None, date_to=None, **filters):
        if set(filters) - QUERY_FIELDS:
            raise ValueError("Unsupported structured history filter")
        matches = []
        for record in self.store.all(self.tenant, "history_record"):
            data = record["data"]
            if any(str(data.get(key, "")).casefold() != str(value).casefold() for key, value in filters.items()):
                continue
            when = data.get("pickup_date")
            if date_from and (not when or when < str(date_from)):
                continue
            if date_to and (not when or when > str(date_to)):
                continue
            matches.append(record)
        return matches

    def profile(self, kind, **filters):
        rows = self.query(**filters)
        dates = sorted(row["data"]["pickup_date"] for row in rows if row["data"].get("pickup_date"))
        facts = {}
        for field in ("customer", "carrier", "equipment", "origin", "destination",
                      "pickup_facility", "delivery_facility"):
            facts[field + "_counts"] = dict(Counter(row["data"][field] for row in rows if row["data"].get(field)))
        rates = []
        currencies = sorted({r["data"]["currency"] for r in rows if r["data"].get("currency")})
        for currency in currencies:
            for metric in ("customer_revenue", "carrier_pay", "gross_margin"):
                samples = []
                for row in rows:
                    data = row["data"]
                    if data.get("currency") != currency:
                        continue
                    value = data.get(metric)
                    if metric == "gross_margin" and value is None and all(
                            data.get(f) is not None for f in ("customer_revenue", "carrier_pay")):
                        value = Decimal(data["customer_revenue"]) - Decimal(data["carrier_pay"])
                    if value is not None:
                        samples.append((Decimal(value), row))
                numbers = [v for v, _ in samples]
                sample_dates = sorted(r["data"]["pickup_date"] for _, r in samples if r["data"].get("pickup_date"))
                rates.append(RateHistorySummary(metric=metric, currency=currency, sample_size=len(numbers),
                    minimum=min(numbers) if numbers else None, maximum=max(numbers) if numbers else None,
                    median=median(numbers) if numbers else None,
                    average=sum(numbers) / len(numbers) if numbers else None,
                    date_from=sample_dates[0] if sample_dates else None,
                    date_to=sample_dates[-1] if sample_dates else None,
                    source_records=[r["id"] for _, r in samples], calculated_at=utcnow()))
        return DerivedProfile(kind=kind, key=filters, sample_size=len(rows),
            date_from=dates[0] if dates else None, date_to=dates[-1] if dates else None,
            source_records=[row["id"] for row in rows], calculated_at=utcnow(), facts=facts, rates=rates)

    def propose_customer_rule(self, customer, phrase):
        if not phrase.strip() or len(phrase) > 200:
            raise ValueError("Bounded source phrase required")
        rows = [row for row in self.query(customer=customer)
                if phrase.casefold() in (row["data"].get("notes") or "").casefold()]
        if len(rows) < 3:
            raise ValueError("At least three distinct historical loads required; one occurrence is not a rule")
        dates = sorted(r["data"]["pickup_date"] for r in rows if r["data"].get("pickup_date"))
        rule = ProposedCustomerRule(id=digest([customer, phrase, [r["id"] for r in rows]]),
            description=phrase, customer=customer, source_records=[r["id"] for r in rows], sample_size=len(rows),
            date_from=dates[0] if dates else None, date_to=dates[-1] if dates else None)
        with self.store.transaction():
            self.store.put(self.tenant, "proposed_customer_rule", rule.id, rule)
        return rule

    def approve_customer_rule(self, rule_id, actor):
        if actor.tenant_id != self.tenant or actor.role != Role.OWNER:
            raise PermissionError("Only the owner can approve customer SOP rules")
        value = self.store.get(self.tenant, "proposed_customer_rule", rule_id)
        value.update(status="APPROVED", approved_by=actor.id, approved_at=utcnow())
        approved = ApprovedCustomerRule.model_validate(value)
        with self.store.transaction():
            self.store.put(self.tenant, "approved_customer_rule", rule_id, approved)
        return approved
